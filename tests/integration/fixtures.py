from pathlib import Path

import pymupdf
from docx import Document as DocxDocument


def make_pdf(path: Path, pages: list[str]) -> Path:
    doc = pymupdf.open()
    try:
        for text in pages:
            page = doc.new_page()
            page.insert_text((72, 72), text, fontsize=11)
        doc.save(path)
    finally:
        doc.close()
    return path


def make_docx(path: Path, sections: list[tuple[str, str]]) -> Path:
    """sections: list of (heading, body) pairs."""
    doc = DocxDocument()
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.save(path)
    return path
